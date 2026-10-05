"""Standalone bounded attachment parser; no application/database/credential imports."""
import json
import os
import subprocess
import sys
import zipfile
from pathlib import Path


def parse(path):
    suffix = path.suffix.lower()
    if suffix == ".txt":
        return [path.read_text(encoding="utf-8-sig")], []
    if suffix == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(path)
        if reader.is_encrypted:
            return [], ["PDF_ENCRYPTED"]
        if len(reader.pages) > 200:
            raise ValueError("too many pages")
        pages = [p.extract_text() or "" for p in reader.pages]
        return pages, [f"page:{i+1}:NO_TEXT_LAYER" for i, p in enumerate(pages) if not p.strip()]
    if suffix == ".docx":
        with zipfile.ZipFile(path) as z:
            entries = z.infolist()
            if len(entries) > 2000 or sum(e.file_size for e in entries) > 60 * 1024 * 1024:
                raise ValueError("decompression budget")
            if any("vbaProject" in e.filename for e in entries):
                raise ValueError("macros not supported")
        from docx import Document
        doc = Document(path)
        parts = [p.text for p in doc.paragraphs]
        parts += [" | ".join(c.text for c in row.cells) for table in doc.tables for row in table.rows]
        return ["\n\n".join(parts)], []
    from PIL import Image
    Image.MAX_IMAGE_PIXELS = 30_000_000
    with Image.open(path) as image:
        image.verify()
    languages = subprocess.run(["tesseract", "--list-langs"], capture_output=True, timeout=5).stdout.decode()
    if "chi_sim" not in languages:
        return [], ["OCR_CHINESE_LANGUAGE_MISSING"]
    result = subprocess.run(["tesseract", str(path), "stdout", "-l", "chi_sim+eng"],
                            capture_output=True, timeout=20, check=True)
    return [result.stdout.decode("utf-8")], ["OCR_UNVERIFIED"]


if __name__ == "__main__":
    if os.name != "nt":
        import resource
        resource.setrlimit(resource.RLIMIT_CPU, (25, 25))
        resource.setrlimit(resource.RLIMIT_AS, (768 * 1024**2, 768 * 1024**2))
        resource.setrlimit(resource.RLIMIT_FSIZE, (25 * 1024**2, 25 * 1024**2))
    try:
        pages, issues = parse(Path(sys.argv[1]))
        text = "\n\n".join(pages)
        if not text.strip():
            result = {"error": "OCR_FAILED", "message": "未得到可读正文；扫描 PDF 或图片需要本地中文 OCR。"}
        else:
            result = {"text": text, "pages": pages, "issues": issues,
                      "completeness": "partial" if issues else "full_text"}
        print(json.dumps(result, ensure_ascii=False))
    except Exception:
        print(json.dumps({"error": "PARSER_FAILED", "message": "文件解析失败，请检查格式或改用正文。"}))
